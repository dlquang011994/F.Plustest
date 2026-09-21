#!/usr/bin/env python3
"""Build the small JSON file consumed by the Dashboard from the source XLSX."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

import openpyxl


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def row_value(row: tuple[object, ...], index: int | None) -> str:
    return text(row[index]) if index is not None and index < len(row) else ""


def rows_with_headers(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        name = text(header)
        if name and name not in indexes:
            indexes[name] = index
    return rows, indexes


def week_number(value: str) -> int:
    digits = "".join(char for char in value if char.isdigit())
    return int(digits) if digits else 0


def region_stats():
    return {
        "nhan_all_w": defaultdict(int),
        "ban_svt_w": defaultdict(int),
        "nhan_all_m": defaultdict(int),
        "ban_svt_m": defaultdict(int),
    }


def build(input_path: Path) -> dict:
    workbook = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    required = {"Volume phiếu nhận", "Volume bán"}
    missing = required - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"Thiếu sheet: {', '.join(sorted(missing))}")

    # Danhba is retained for ASM/shop metadata. It must not override the sales region.
    shop_to_region: dict[str, str] = {}
    shop_to_asm: dict[str, str] = {}
    shop_to_original: dict[str, str] = {}
    asm_original: dict[str, str] = {}
    all_asms: set[str] = set()
    all_regions: set[str] = set()

    if "Danhba" in workbook.sheetnames:
        rows, indexes = rows_with_headers(workbook["Danhba"])
        for row in rows:
            shop = row_value(row, indexes.get("Cửa hàng"))
            region = row_value(row, indexes.get("Vùng miền"))
            asm = row_value(row, indexes.get("Email ASM"))
            if not shop:
                continue
            key = shop.lower()
            shop_to_original[key] = shop
            if region and region.lower() != "loại":
                shop_to_region[key] = region
                all_regions.add(region)
            if asm:
                asm_key = asm.lower()
                shop_to_asm[key] = asm_key
                asm_original[asm_key] = asm
                all_asms.add(asm)

    satellite_shops: set[str] = set()
    if "Shop vệ tinh" in workbook.sheetnames:
        rows, indexes = rows_with_headers(workbook["Shop vệ tinh"])
        for row in rows:
            shop = row_value(row, indexes.get("Shop") if "Shop" in indexes else indexes.get("Cửa hàng"))
            if shop:
                satellite_shops.add(shop.lower())

    receive_rows, receive = rows_with_headers(workbook["Volume phiếu nhận"])
    sales_rows, sales = rows_with_headers(workbook["Volume bán"])

    # First pass: establish available weeks/months and fallback shop metadata from receipts.
    receive_data: list[tuple[str, str, str, str]] = []
    weeks: set[str] = set()
    months: set[str] = set()
    for row in receive_rows:
        shop = row_value(row, receive.get("Cửa hàng"))
        region = row_value(row, receive.get("Vùng kinh doanh shop"))
        week = row_value(row, receive.get("Tuần")).upper()
        month = row_value(row, receive.get("Tháng"))
        receive_data.append((shop, region, week, month))
        if week.startswith("W"):
            weeks.add(week)
        if month:
            months.add(month)
        if shop:
            key = shop.lower()
            shop_to_original[key] = shop
            if region and region.lower() != "loại" and key not in shop_to_region:
                shop_to_region[key] = region
        if region and region.lower() != "loại":
            all_regions.add(region)

    sales_data: list[tuple[str, str, str, str, str]] = []
    for row in sales_rows:
        shop = row_value(row, sales.get("Shop Name") if "Shop Name" in sales else sales.get("Cửa hàng"))
        classification = row_value(row, sales.get("Phân loại") if "Phân loại" in sales else sales.get("Phân Loại"))
        # Use the sales row region as the source of truth for sales counts.
        region = row_value(row, sales.get("Vùng kinh doạnh") if "Vùng kinh doạnh" in sales else sales.get("Vùng kinh doanh"))
        week = row_value(row, sales.get("Tuần")).upper()
        month = row_value(row, sales.get("Tháng"))
        sales_data.append((shop, classification, region, week, month))
        if week.startswith("W"):
            weeks.add(week)
        if month:
            months.add(month)
        if region and region.lower() != "loại":
            all_regions.add(region)

    all_weeks = sorted(weeks, key=week_number)
    latest_weeks = all_weeks[-4:]
    current_week = latest_weeks[-1] if latest_weeks else ""
    previous_week = latest_weeks[-2] if len(latest_weeks) > 1 else ""
    all_months = sorted(months, key=week_number)
    months_to_show = [month for month in all_months if week_number(month) >= 9]

    stats: dict[str, dict] = defaultdict(region_stats)
    shop_stats: dict[str, dict] = {}
    asm_stats: dict[str, dict] = {}

    def shop_stat(key: str, fallback: str) -> dict:
        if key not in shop_stats:
            asm_key = shop_to_asm.get(key)
            shop_stats[key] = {
                "shopName": shop_to_original.get(key, fallback or key),
                "vung": shop_to_region.get(key, "Khác"),
                "asm": asm_original.get(asm_key, "") if asm_key else "",
                "asmKey": asm_key,
                "svt": key in satellite_shops,
                "nhan_all_w": defaultdict(int),
                "ban_svt_cur": 0,
            }
        return shop_stats[key]

    def asm_stat(key: str) -> dict:
        if key not in asm_stats:
            region = next((shop_to_region.get(shop_key, "") for shop_key, asm_key in shop_to_asm.items() if asm_key == key), "")
            asm_stats[key] = {
                "asmName": asm_original.get(key, key),
                "vung": region or "Khác",
                "nhan_all_w": defaultdict(int),
                "ban_svt_cur": 0,
            }
        return asm_stats[key]

    receipt_current = receipt_previous = 0
    for shop, raw_region, week, month in receive_data:
        if not shop:
            continue
        key = shop.lower()
        region = shop_to_region.get(key) or raw_region
        if not region or region.lower() == "loại":
            continue
        if month in months_to_show:
            stats[region]["nhan_all_m"][month] += 1
        if week in latest_weeks:
            stats[region]["nhan_all_w"][week] += 1
            if week == current_week:
                receipt_current += 1
            if week == previous_week:
                receipt_previous += 1
            shop_stat(key, shop)["nhan_all_w"][week] += 1
            asm_key = shop_to_asm.get(key)
            if asm_key:
                asm_stat(asm_key)["nhan_all_w"][week] += 1

    sales_current = sales_previous = 0
    for shop, classification, region, week, month in sales_data:
        if not shop or classification.lower() != "lấy":
            continue
        if not region or region.lower() == "loại":
            continue
        key = shop.lower()
        if month in months_to_show:
            stats[region]["ban_svt_m"][month] += 1
        if week in latest_weeks:
            stats[region]["ban_svt_w"][week] += 1
            if week == current_week:
                sales_current += 1
                shop_stat(key, shop)["ban_svt_cur"] += 1
                asm_key = shop_to_asm.get(key)
                if asm_key:
                    asm_stat(asm_key)["ban_svt_cur"] += 1
            if week == previous_week:
                sales_previous += 1

    regions = sorted(region for region in stats if region != "Khác" and region in all_regions)
    main_rows, vuong_rows, luyke_rows = [], [], []
    main_total = {"vùng": "Tổng", "tổng": 0, "sl_tang_giam": 0}
    vuong_total = {"vùng": "Tổng", "tổng": 0, "ban_w_cur": 0}
    luyke_total = {"vùng": "Tổng", "tang_giam": 0}
    for week in latest_weeks:
        main_total[week.lower()] = 0
        vuong_total[week.lower()] = 0
    for month in months_to_show:
        luyke_total[f"Th{month}"] = 0
    total_receipts_by_month = defaultdict(int)
    total_sales_by_month = defaultdict(int)

    for region in regions:
        stat = stats[region]
        main = {"vùng": region, "tổng": 0}
        vuong = {"vùng": region, "tổng": 0}
        for week in latest_weeks:
            key = week.lower()
            value = stat["nhan_all_w"][week]
            main[key] = value
            main["tổng"] += value
            main_total[key] += value
            main_total["tổng"] += value
            vuong[key] = value
            vuong[f"ban_{key}"] = stat["ban_svt_w"][week]
            vuong["tổng"] += value
            vuong_total[key] += value
            vuong_total["tổng"] += value
        current_receipts = stat["nhan_all_w"][current_week]
        previous_receipts = stat["nhan_all_w"][previous_week]
        main["sl_tang_giam"] = current_receipts - previous_receipts
        main["ti_le_tang_giam"] = ((current_receipts - previous_receipts) / previous_receipts) if previous_receipts else (1 if current_receipts else 0)
        main_total["sl_tang_giam"] += main["sl_tang_giam"]
        vuong["ban_w_cur"] = stat["ban_svt_w"][current_week]
        vuong["ti_le"] = current_receipts / vuong["ban_w_cur"] if vuong["ban_w_cur"] else 0
        vuong_total["ban_w_cur"] += vuong["ban_w_cur"]
        main_rows.append(main)
        vuong_rows.append(vuong)

        cumulative = {"vùng": region}
        for month in months_to_show:
            receipts = stat["nhan_all_m"][month]
            sales = stat["ban_svt_m"][month]
            cumulative[f"Th{month}"] = receipts / sales if sales else 0
            total_receipts_by_month[month] += receipts
            total_sales_by_month[month] += sales
        luyke_rows.append(cumulative)

    current_month = months_to_show[-1] if months_to_show else None
    previous_month = months_to_show[-2] if len(months_to_show) > 1 else None
    for row in luyke_rows:
        row["tang_giam"] = (row.get(f"Th{current_month}", 0) - row.get(f"Th{previous_month}", 0)) if current_month else 0
    for month in months_to_show:
        luyke_total[f"Th{month}"] = total_receipts_by_month[month] / total_sales_by_month[month] if total_sales_by_month[month] else 0
    luyke_total["tang_giam"] = (luyke_total.get(f"Th{current_month}", 0) - luyke_total.get(f"Th{previous_month}", 0)) if current_month else 0

    total_current = main_total.get(current_week.lower(), 0)
    total_previous = main_total.get(previous_week.lower(), 0)
    main_total["ti_le_tang_giam"] = ((total_current - total_previous) / total_previous) if total_previous else (1 if total_current else 0)
    vuong_total["ti_le"] = vuong_total.get(current_week.lower(), 0) / vuong_total["ban_w_cur"] if vuong_total["ban_w_cur"] else 0

    asm_rows = []
    for asm_key, stat in asm_stats.items():
        current_receipts = stat["nhan_all_w"][current_week]
        previous_receipts = stat["nhan_all_w"][previous_week]
        asm_rows.append({
            "asm": stat["asmName"], "asmKey": asm_key, "vùng": stat["vung"],
            "ban_w_cur": stat["ban_svt_cur"], previous_week.lower(): previous_receipts,
            current_week.lower(): current_receipts, "tang_giam": current_receipts - previous_receipts,
            "ti_le": current_receipts / stat["ban_svt_cur"] if stat["ban_svt_cur"] else 0,
        })

    shop_rows = []
    for stat in shop_stats.values():
        row = {"shop": stat["shopName"], "asm": stat["asm"], "vùng": stat["vung"], "svt": stat["svt"], "tổng": 0, "ban_w_cur": stat["ban_svt_cur"]}
        for week in latest_weeks:
            value = stat["nhan_all_w"][week]
            row[week.lower()] = value
            row["tổng"] += value
        row["ti_le"] = row.get(current_week.lower(), 0) / row["ban_w_cur"] if row["ban_w_cur"] else 0
        shop_rows.append(row)

    return {
        "weeks": latest_weeks, "months": months_to_show, "w_cur": current_week, "w_prv": previous_week,
        "kpi": {
            "n_cur": receipt_current, "n_prv": receipt_previous, "wow_n": receipt_current - receipt_previous,
            "ban_svt_cur": sales_current, "ban_svt_prv": sales_previous,
            "ti_le_cur": receipt_current / sales_current if sales_current else 0,
            "ti_le_prv": receipt_previous / sales_previous if sales_previous else 0,
            "wow_ti_le": (receipt_current / sales_current if sales_current else 0) - (receipt_previous / sales_previous if sales_previous else 0),
            "w_cur": current_week, "w_prv": previous_week,
        },
        "vùng_list": sorted(region for region in all_regions if region != "Khác"), "asm_list": sorted(all_asms),
        "main_rows": main_rows, "main_tot": main_total, "luyke_rows": luyke_rows, "luyke_tot": luyke_total,
        "vuong_rows": vuong_rows, "vuong_tot": vuong_total, "asm_rows": sorted(asm_rows, key=lambda row: row["asm"]),
        "shop_rows": sorted(shop_rows, key=lambda row: row["shop"]),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    data = build(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
